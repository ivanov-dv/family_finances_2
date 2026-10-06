import io
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from transactions.models import Transaction

pytestmark = pytest.mark.django_db(transaction=True)


@pytest.fixture
def make_tx(user_1, owner_space):
    core = user_1.core_settings

    def _make(group='Продукты', description='', value='10.00', created_at=None):
        tx = Transaction.objects.create(
            space=owner_space, author=user_1,
            period_month=core.current_month, period_year=core.current_year,
            type_transaction='expense', group_name=group, description=description,
            value_transaction=Decimal(value),
        )
        if created_at:
            Transaction.objects.filter(pk=tx.pk).update(created_at=created_at)
        return tx

    return _make


def export(client):
    response = client.get(reverse('export:excel'))
    assert response.status_code == 200
    return response.content


def rows(content):
    return list(load_workbook(io.BytesIO(content)).active.iter_rows(min_row=2))


class TestExportContent:

    def test_user_text_never_becomes_a_formula(self, user_1_client, make_tx):
        make_tx(group='=1+1', description='=HYPERLINK("http://example.invalid","x")')
        content = export(user_1_client)
        row = rows(content)[0]
        assert (row[2].data_type, row[2].value) == ('s', '=1+1')
        assert row[4].data_type == 's'
        assert row[4].value == '=HYPERLINK("http://example.invalid","x")'
        sheet_xml = zipfile.ZipFile(io.BytesIO(content)).read('xl/worksheets/sheet1.xml').decode()
        assert '<f>' not in sheet_xml

    @pytest.mark.parametrize('description', ['=2+2', '+7 999', '-5 руб', '@SUM(1,2)'])
    def test_formula_like_text_gets_quote_prefix(self, user_1_client, make_tx, description):
        make_tx(description=description)
        cell = rows(export(user_1_client))[0][4]
        assert cell.value == description
        assert cell.quotePrefix is True

    def test_plain_text_has_no_quote_prefix(self, user_1_client, make_tx):
        make_tx(description='обычный текст')
        assert rows(export(user_1_client))[0][4].quotePrefix is False

    def test_control_characters_do_not_break_the_export(self, user_1_client, make_tx):
        """Раньше один символ U+000B в описании делал выгрузку периода недоступной (500)."""
        make_tx(description='аб\x0bвг\x01д')
        assert rows(export(user_1_client))[0][4].value == 'абвгд'

    def test_date_is_local_and_a_real_date(self, user_1_client, make_tx):
        """21:30 UTC — уже 00:30 следующего дня по Москве (TIME_ZONE проекта)."""
        make_tx(created_at=datetime(2026, 10, 5, 21, 30, tzinfo=timezone.utc))
        cell = rows(export(user_1_client))[0][0]
        assert cell.value.date() == date(2026, 10, 6)
        assert cell.number_format == 'DD.MM.YYYY'

    def test_amount_stays_a_number(self, user_1_client, make_tx):
        make_tx(value='1234.56')
        assert rows(export(user_1_client))[0][3].value == pytest.approx(1234.56)
